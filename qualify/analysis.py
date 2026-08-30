"""Can the camera certify the part? Three measured chapters on open data.

A. Probability of detection (POD) of layer-image anomalies, from the Aalto
   University LPBF dataset (EOS M290, 316L; powder-bed camera 1280x1024 and
   optical tomography 2000x2000; LabelImg boxes; CC BY 4.0, Zenodo 14996806).
   A transparent detector is trained on two builds and evaluated on the third,
   the operating point is set by a false-call budget on defect-free images,
   and the hit/miss results are fitted with the MIL-HDBK-1823A logistic model
   POD(a) = logistic(b0 + b1 ln a) to give a50, a90 and a90/95.
B. Witness-coupon statistics, from the PSU Ti-6Al-4V process-structure-property
   dataset (42 parameter sets, 168 tensile specimens; CC BY 4.0, Zenodo 6587905):
   a process-control reference distribution, operating-characteristic curves
   for an n-coupon acceptance rule, and strength against porosity with the
   DISCMAM density gate drawn on it.
C. The four-tier release policy from the console, applied to the 42 real process
   windows: density from X-ray CT, witness result against the ASTM F2924 minimums.

Everything written to docs/qualify/data.json; the page reads that file.
"""
import glob
import json
import math
import os
import sys
import time
import xml.etree.ElementTree as ET

import numpy as np
from PIL import Image
from scipy import ndimage
from scipy.optimize import brentq
from scipy.stats import norm

ROOT = os.path.expanduser("~/projects/am-datasets/aalto/Original Images and manual labels")
PSU = os.path.expanduser("~/projects/am-datasets/psu-ti64")
FATIGUE = os.path.expanduser("~/projects/am-datasets/fatigue/FatigueData-AM2022.xlsx")
OUT = "docs/qualify/data.json"

SENSORS = {
    # the powder-bed camera sees recoater streaks: thin horizontal lines, a few pixels tall, tens to hundreds of pixels long
    "PB": {"label": "powder-bed camera", "mm_per_px": 250 / 1100, "cls": "defects", "roi": (210, 880, 150, 1130), "win": (12, 96), "stride": (8, 48),
           "what": "recoater streaks: thin horizontal lines in the fresh powder layer"},
    # optical tomography is a colour-mapped heat image: blue background, light-blue exposures, red and white where the melt overheated
    "OT": {"label": "optical tomography", "mm_per_px": 250 / 1920, "cls": "overheated", "roi": (40, 1960, 40, 1960), "win": (24, 24), "stride": (12, 12),
           "what": "overheated melt: red and white areas inside the exposed cross-sections of the colour-mapped heat image"},
}
TEST_BUILD = "SI383820240318120348"   # the largest build is held out
OPERATING_POINTS = [0.25, 1, 3, 10]   # accepted false windows per inspected image; the page lets the reader move this
PRIMARY_OP = 3
FEATURE_NAMES = {"PB": ["line response averaged over 96 px", "over 200 px", "over 400 px", "ridge response over 200 px"],
                 "OT": ["redness, max", "redness, mean", "share of hot pixels", "red channel, max"]}


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


# ---------------------------------------------------------------- A. images
def load_image(path, sensor):
    im = Image.open(path)
    if sensor == "PB":
        return np.asarray(im.convert("L"), dtype=np.float32)
    return np.asarray(im.convert("RGB"), dtype=np.float32)


def grid(shape, cfg):
    y0, y1, x0, x1 = cfg["roi"]
    wh, ww = cfg["win"]; sh, sw = cfg["stride"]
    ys = np.arange(y0 + wh // 2, y1 - wh // 2, sh)
    xs = np.arange(x0 + ww // 2, x1 - ww // 2, sw)
    return np.meshgrid(ys, xs, indexing="ij")


def window_features(img, cfg, sensor):
    """Per-window features on a stride grid inside the region of interest (the build plate)."""
    wh, ww = cfg["win"]
    if sensor == "PB":
        hp = img - ndimage.uniform_filter1d(img, 21, axis=0)          # thin horizontal structure survives, slow shading does not
        y0, y1, x0, x1 = cfg["roi"]
        noise = np.median(np.abs(hp[y0:y1, x0:x1])) + 1e-3
        hpn = hp / noise
        ridge = hpn - 0.5 * (np.roll(hpn, 6, 0) + np.roll(hpn, -6, 0))   # a line is brighter or darker than the rows six pixels away
        F = [ndimage.maximum_filter(np.abs(ndimage.uniform_filter1d(hpn, L, axis=1)), (wh, 1)) for L in (96, 200, 400)]
        F.append(ndimage.maximum_filter(np.abs(ndimage.uniform_filter1d(ridge, 200, axis=1)), (wh, 1)))
    else:
        red = img[..., 0] - img[..., 2]                                 # red minus blue: negative on the blue background and the light-blue exposures
        hot = (red > 30).astype(np.float32)
        F = [ndimage.maximum_filter(red, (wh, ww)), ndimage.uniform_filter(red, (wh, ww)), ndimage.uniform_filter(hot, (wh, ww)), ndimage.maximum_filter(img[..., 0], (wh, ww))]
    Y, X = grid(img.shape[:2], cfg)
    return np.stack([f[Y, X] for f in F], axis=-1), Y, X


def boxes_for(sensor, name):
    p = f"{ROOT}/{sensor}_label/{name}.xml"
    if not os.path.exists(p):
        return []
    t = ET.parse(p).getroot()
    out = []
    for o in t.findall("object"):
        b = o.find("bndbox")
        out.append([int(b.find(k).text) for k in ("xmin", "ymin", "xmax", "ymax")])
    return out


def fit_logreg(X, y, iters=300, l2=1e-3):
    """Plain logistic regression by Newton's method on standardised features."""
    mu, sd = X.mean(0), X.std(0) + 1e-6
    Z = np.hstack([np.ones((len(X), 1)), (X - mu) / sd])
    w = np.zeros(Z.shape[1])
    for _ in range(iters):
        p = 1 / (1 + np.exp(-Z @ w))
        g = Z.T @ (p - y) + l2 * w
        H = (Z * (p * (1 - p))[:, None]).T @ Z + l2 * np.eye(len(w))
        step = np.linalg.solve(H, g)
        w -= step
        if np.abs(step).max() < 1e-6:
            break
    return {"w": w, "mu": mu, "sd": sd}


def predict(model, X):
    Z = np.hstack([np.ones((len(X), 1)), (X - model["mu"]) / model["sd"]])
    return 1 / (1 + np.exp(-Z @ model["w"]))


def pod_fit(sizes, hits):
    """MIL-HDBK-1823A hit/miss model: logit POD = b0 + b1 ln a; Wald-type 95% lower bound."""
    a = np.log(np.asarray(sizes, float))
    y = np.asarray(hits, float)
    Z = np.stack([np.ones_like(a), a], 1)
    w = np.zeros(2)
    for _ in range(200):
        p = 1 / (1 + np.exp(-Z @ w))
        g = Z.T @ (p - y) + 1e-4 * w
        H = (Z * (p * (1 - p))[:, None]).T @ Z + 1e-4 * np.eye(2)
        step = np.linalg.solve(H, g)
        w -= step
        if np.abs(step).max() < 1e-8:
            break
    p = 1 / (1 + np.exp(-Z @ w))
    cov = np.linalg.inv((Z * (p * (1 - p))[:, None]).T @ Z + 1e-4 * np.eye(2))
    b0, b1 = w

    def sig(z):
        return 1 / (1 + math.exp(-max(-500, min(500, z))))

    def pod(x):
        return sig(b0 + b1 * math.log(x))

    def pod_lower(x):
        la = math.log(x)
        eta = b0 + b1 * la
        se = math.sqrt(max(cov[0, 0] + 2 * la * cov[0, 1] + la * la * cov[1, 1], 0))
        return sig(eta - 1.645 * se)

    def solve(f, target, lo, hi):
        try:
            return brentq(lambda x: f(x) - target, lo, hi)
        except ValueError:
            return None
    lo, hi = float(np.exp(a.min())) / 4, float(np.exp(a.max())) * 20
    return {"b0": float(b0), "b1": float(b1), "cov": cov.tolist(),
            "a50": solve(pod, 0.5, lo, hi), "a90": solve(pod, 0.9, lo, hi), "a90_95": solve(pod_lower, 0.9, lo, hi),
            "curve": [[float(x), pod(x), pod_lower(x)] for x in np.exp(np.linspace(a.min() - 0.5, a.max() + 0.5, 60))]}


def chapter_a():
    out = {}
    for sensor, cfg in SENSORS.items():
        log(f"[{sensor}] loading")
        imgs = sorted(glob.glob(f"{ROOT}/{sensor}/*.jpg"))
        if "--quick" in sys.argv:
            rng = np.random.default_rng(3); imgs = [imgs[i] for i in sorted(rng.choice(len(imgs), 320, replace=False))]
        train_feats, train_lab = [], []
        test = []  # (name, F, Y, X, boxes, near)
        n_train = n_test = n_unl = 0
        wh, ww = cfg["win"]
        for path in imgs:
            name = os.path.basename(path)[:-4]
            build = name.split("_")[0]
            boxes = boxes_for(sensor, name)
            if not boxes:
                n_unl += 1
                continue  # only images an annotator inspected and boxed are used; unlabelled images are not assumed clean
            img = load_image(path, sensor)
            F, Y, X = window_features(img, cfg, sensor)
            inside = np.zeros(Y.shape, bool)
            near = np.zeros(Y.shape, bool)
            for x0, y0, x1, y1 in boxes:
                inside |= (X >= x0) & (X <= x1) & (Y >= y0 - wh // 2) & (Y <= y1 + wh // 2)
                near |= (X >= x0 - 2 * ww) & (X <= x1 + 2 * ww) & (Y >= y0 - 3 * wh) & (Y <= y1 + 3 * wh)
            if build == TEST_BUILD:
                test.append((name, F, Y, X, boxes, near))
                n_test += 1
            else:
                pos = F[inside]
                neg = F[~near]
                if len(neg) > 400:
                    idx = np.random.default_rng(len(train_lab)).choice(len(neg), 400, replace=False)
                    neg = neg[idx]
                train_feats.append(pos); train_lab.append(np.ones(len(pos)))
                train_feats.append(neg); train_lab.append(np.zeros(len(neg)))
                n_train += 1
        Xtr = np.vstack(train_feats); ytr = np.concatenate(train_lab)
        log(f"[{sensor}] train images {n_train}, test images {n_test}, unlabelled left aside {n_unl}, windows {len(ytr)} ({int(ytr.sum())} inside boxes)")
        model = fit_logreg(Xtr, ytr)
        # the clean reference: windows of the held-out images that lie away from every box (the annotator looked and drew nothing there)
        neg_scores = []
        box_records = []  # size_mm, score, length_px
        mm = cfg["mm_per_px"]
        for name, F, Y, X, boxes, near in test:
            S = predict(model, F.reshape(-1, F.shape[-1])).reshape(F.shape[:2])
            neg_scores.append(S[~near])
            for x0, y0, x1, y1 in boxes:
                sel = (X >= x0 - ww / 2) & (X <= x1 + ww / 2) & (Y >= y0 - wh) & (Y <= y1 + wh)
                if not sel.any():
                    continue
                size_px = max(x1 - x0, y1 - y0)
                if size_px < 2:
                    continue
                box_records.append((size_px * mm, float(S[sel].max())))
        neg = np.concatenate(neg_scores)
        windows_per_img = len(neg) / n_test
        sizes = np.array([b[0] for b in box_records]); scores = np.array([b[1] for b in box_records])
        ops = {}
        for fc in OPERATING_POINTS:
            thr = float(np.quantile(neg, 1 - fc / windows_per_img))
            hits = scores > thr
            fit = pod_fit(sizes, hits) if 0.02 < hits.mean() < 0.98 else None
            edges = np.quantile(sizes, np.linspace(0, 1, 9))
            bins = []
            for i in range(8):
                sel = (sizes >= edges[i]) & ((sizes <= edges[i + 1]) if i == 7 else (sizes < edges[i + 1]))
                if sel.sum():
                    n, k = int(sel.sum()), int(hits[sel].sum())
                    # Clopper-Pearson 95% lower bound on the hit rate in the bin
                    from scipy.stats import beta
                    lo = float(beta.ppf(0.05, k, n - k + 1)) if k > 0 else 0.0
                    bins.append({"lo": float(edges[i]), "hi": float(edges[i + 1]), "n": n, "hit": k / n, "lower95": lo, "mid": float(np.exp(np.log(sizes[sel]).mean()))})
            ops[str(fc)] = {"threshold": thr, "hit_rate": float(hits.mean()), "n": int(len(hits)), "fit": fit, "bins": bins}
            log(f"[{sensor}] {fc} false calls/image: hit rate {hits.mean():.3f}" + (f", a50 {fit['a50']}, a90 {fit['a90']}, a90/95 {fit['a90_95']}" if fit else " (saturated, binomial bounds only)"))
        qs = np.linspace(0, 1, 2001)
        tail = np.sort(neg)[::-1][:25000]   # the highest reference scores, so the page can set any false-call budget exactly
        out[sensor] = {"label": cfg["label"], "cls": cfg["cls"], "what": cfg["what"], "mm_per_px": mm, "roi": cfg["roi"],
                       "mm_note": f"assumed: the 250 mm plate spans about {round(250 / mm)} px of the image, read off the plate edges; an estimate, not a calibration",
                       "train_images": n_train, "test_images": n_test, "unlabelled_images": n_unl, "windows_per_image": windows_per_img,
                       "boxes": int(len(sizes)), "size_p10": float(np.quantile(sizes, 0.1)), "size_p50": float(np.quantile(sizes, 0.5)), "size_p90": float(np.quantile(sizes, 0.9)),
                       "ops": ops, "primary": str(PRIMARY_OP), "weights": model["w"].tolist(), "features": FEATURE_NAMES[sensor],
                       "neg_quantiles": [float(v) for v in np.quantile(neg, qs)], "neg_tail": [round(float(v), 5) for v in tail], "neg_total": int(len(neg)), "box_scores": [[round(float(a), 3), round(float(b), 5)] for a, b in box_records]}
    return out


# ---------------------------------------------------------------- B. coupons
def read_psu():
    import openpyxl
    wb = openpyxl.load_workbook(f"{PSU}/psp_feature_table.xlsx", read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    sets = []
    for r in rows[6:]:
        if r[1] is None or not str(r[1]).strip().isdigit():
            continue
        if any(r[i] is None for i in (2, 3, 4, 7, 9, 11, 33, 34, 35, 36, 39, 40)):
            continue  # a parameter set without a full record is left out, and counted
        sets.append({"set": int(r[1]), "power_w": float(r[2]), "speed_mm_s": float(r[3]), "led_j_m": float(r[4]),
                     "porosity_xct_pct": float(r[7]), "porosity_arch_pct": float(r[9]), "pores": int(float(r[11])),
                     "uts_mean": float(r[33]), "uts_sd": float(r[34]), "ys_mean": float(r[35]), "ys_sd": float(r[36]),
                     "el_mean": float(r[39]), "el_sd": float(r[40])})
    wb2 = openpyxl.load_workbook(f"{PSU}/mechanical_property_table.xlsx", read_only=True, data_only=True)
    specimens = []
    for ws in wb2.worksheets:
        for r in list(ws.iter_rows(values_only=True))[2:]:
            if r[0] is None or not str(r[0]).strip().isdigit():
                continue
            try:
                specimens.append({"set": int(r[0]), "sheet": ws.title, "uts": float(r[1]), "ys": float(r[2]), "el": float(r[4])})
            except (TypeError, ValueError):
                pass
    return sets, specimens


F2924 = {"uts": 895, "ys": 825, "el": 10}   # ASTM F2924 minimums for Ti-6Al-4V powder-bed fusion parts


def chapter_b(sets, specimens):
    by_set = {}
    for s in specimens:
        by_set.setdefault(s["set"], []).append(s)
    # the reference population: specimens from parameter sets that pass the density gate
    dense = {s["set"] for s in sets if 100 - s["porosity_xct_pct"] >= 99.9}
    ref = [s for s in specimens if s["set"] in dense]
    uts = np.array([s["uts"] for s in ref]); el = np.array([s["el"] for s in ref])
    pcrd = {"n": len(ref), "sets": len(dense), "uts_mean": float(uts.mean()), "uts_sd": float(uts.std(ddof=1)), "el_mean": float(el.mean()), "el_sd": float(el.std(ddof=1)),
            "uts_p1": float(np.quantile(uts, 0.01)), "el_p1": float(np.quantile(el, 0.01))}
    # OC curves: accept a build if all n coupons exceed the limit L = mean - k sd; probability of acceptance when the
    # true mean has dropped by delta percent (same sd). k chosen so that a healthy build passes 6 coupons 95% of the time.
    sd = pcrd["uts_sd"]; mu = pcrd["uts_mean"]
    def p_accept(n, k, shift_pct):
        L = mu - k * sd
        mu2 = mu * (1 - shift_pct / 100)
        p1 = 1 - norm.cdf((L - mu2) / sd)
        return p1 ** n
    k = brentq(lambda kk: p_accept(6, kk, 0) - 0.95, 0.5, 6)
    shifts = list(np.linspace(0, 12, 25))
    oc = {str(n): [p_accept(n, k, s) for s in shifts] for n in (1, 2, 4, 6, 12)}
    # how many coupons to catch a 5% mean drop with 90% probability
    need = {}
    for drop in (3, 5, 8):
        n90 = None
        for n in range(1, 60):
            if 1 - p_accept(n, k, drop) >= 0.9:
                n90 = n; break
        need[str(drop)] = n90
    # strength vs porosity per set
    scatter = [{"set": s["set"], "porosity": s["porosity_xct_pct"], "density": 100 - s["porosity_xct_pct"], "uts": s["uts_mean"], "el": s["el_mean"], "ys": s["ys_mean"], "led": s["led_j_m"], "power": s["power_w"], "speed": s["speed_mm_s"]} for s in sets]
    por = np.array([s["porosity_xct_pct"] for s in sets]); u = np.array([s["uts_mean"] for s in sets]); e = np.array([s["el_mean"] for s in sets])
    corr_uts = float(np.corrcoef(np.log10(por + 1e-3), u)[0, 1]); corr_el = float(np.corrcoef(np.log10(por + 1e-3), e)[0, 1])
    return {"pcrd": pcrd, "k": float(k), "limit_uts": float(mu - k * sd), "shifts": shifts, "oc": oc, "need": need, "scatter": scatter,
            "corr_log_porosity_uts": corr_uts, "corr_log_porosity_el": corr_el, "f2924": F2924, "n_specimens": len(specimens), "n_sets": len(sets)}


# ---------------------------------------------------------------- C. the policy on real windows
def chapter_c(sets, specimens):
    by_set = {}
    for s in specimens:
        by_set.setdefault(s["set"], []).append(s)
    rows = []
    for s in sets:
        sp = by_set.get(s["set"], [])
        witness = "pass" if sp and all(x["uts"] >= F2924["uts"] and x["ys"] >= F2924["ys"] and x["el"] >= F2924["el"] for x in sp) else ("fail" if sp else "none")
        density = 100 - s["porosity_xct_pct"]
        if density >= 99.9 and witness == "pass":
            tier = "permanent"
        elif density >= 99.5 and witness == "pass":
            tier = "temporary"
        elif density >= 99.0:
            tier = "emergency only"
        else:
            tier = "scrap"
        rows.append({"set": s["set"], "power": s["power_w"], "speed": s["speed_mm_s"], "led": s["led_j_m"], "density": density, "witness": witness, "coupons": len(sp),
                     "uts_min": min(x["uts"] for x in sp) if sp else None, "el_min": min(x["el"] for x in sp) if sp else None, "tier": tier})
    counts = {}
    for r in rows:
        counts[r["tier"]] = counts.get(r["tier"], 0) + 1
    # the interesting cases: dense but failing the witness (density is not strength), and the reverse
    dense_fail = [r["set"] for r in rows if r["density"] >= 99.9 and r["witness"] != "pass"]
    porous_pass = [r["set"] for r in rows if r["density"] < 99.9 and r["witness"] == "pass"]
    return {"rows": rows, "counts": counts, "dense_but_witness_fail": dense_fail, "porous_but_witness_pass": porous_pass}


# ---------------------------------------------------------------- D. fatigue scatter
def chapter_d():
    import openpyxl
    wb = openpyxl.load_workbook(FATIGUE, read_only=True, data_only=True)
    par = list(wb["parameter"].iter_rows(values_only=True))
    hdr = [str(x).replace("\n", " ").strip() if x else "" for x in par[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    def col(name):
        for h, i in idx.items():
            if h.lower().startswith(name):
                return i
        return None
    c_id, c_mat, c_am, c_seq, c_sur = col("dataset id"), col("name of the material"), col("types of am"), col("processing sequence"), col("specimens description")
    keep = {}
    for r in par[2:]:
        if r[c_id] is None:
            continue
        mat = str(r[c_mat] or "").lower(); am = str(r[c_am] or "").lower()
        if ("ti-6al-4v" in mat or "ti6al4v" in mat or "ti64" in mat) and ("pbf-lb" in am or "l-pbf" in am or "slm" in am or "lpbf" in am or "laser powder" in am or "pbf-l" in am):
            keep[str(r[c_id])] = {"seq": str(r[c_seq] or ""), "spec": str(r[c_sur] or "")}
    sn = list(wb["S-N"].iter_rows(values_only=True))
    pts = []
    for r in sn[1:]:
        if r[0] is None or str(r[0]) not in keep:
            continue
        try:
            N, sa, ro = float(r[1]), float(r[2]), int(float(r[3] or 0))
        except (TypeError, ValueError):
            continue
        if N > 0 and sa > 0:
            pts.append([N, sa, ro, str(r[0])])
    # scatter of log-life in stress bins
    A = np.array([[p[0], p[1]] for p in pts if p[2] == 0])
    bins = []
    if len(A):
        edges = np.quantile(A[:, 1], np.linspace(0, 1, 7))
        for i in range(6):
            sel = (A[:, 1] >= edges[i]) & (A[:, 1] <= edges[i + 1])
            if sel.sum() >= 5:
                ln = np.log10(A[sel, 0])
                bins.append({"lo": float(edges[i]), "hi": float(edges[i + 1]), "n": int(sel.sum()), "log_life_sd": float(ln.std(ddof=1)), "log_life_mean": float(ln.mean())})
    return {"datasets": len(keep), "points": len(pts), "runouts": int(sum(p[2] for p in pts)), "bins": bins,
            "sample": [[round(p[0]), round(p[1], 1), p[2]] for p in pts[::max(1, len(pts) // 1500)]]}


if __name__ == "__main__":
    t0 = time.time()
    A = chapter_a()
    if "--quick" in sys.argv:
        sys.exit(0)
    sets, specimens = read_psu()
    B = chapter_b(sets, specimens)
    C = chapter_c(sets, specimens)
    try:
        D = chapter_d()
    except Exception as e:  # noqa: BLE001
        log("fatigue chapter failed:", repr(e)); D = None
    json.dump({"built": "2026-08-30", "seconds": time.time() - t0, "test_build": TEST_BUILD, "pod": A, "coupons": B, "policy": C, "fatigue": D}, open(OUT, "w"))
    log(f"wrote {OUT} in {time.time()-t0:.0f}s")
    log("PCRD", B["pcrd"], "k", round(B["k"], 2), "need", B["need"], "corr", round(B["corr_log_porosity_uts"], 2), round(B["corr_log_porosity_el"], 2))
    log("policy counts", C["counts"], "dense-but-fail", C["dense_but_witness_fail"], "porous-but-pass", C["porous_but_witness_pass"])
    if D: log("fatigue", D["datasets"], "datasets", D["points"], "points", D["bins"][:2])
